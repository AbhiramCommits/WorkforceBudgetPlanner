import xml.etree.ElementTree as ET
import zipfile

import pandas as pd
from openpyxl import load_workbook

from wbp.excel_report import build_report


def _payload() -> dict:
    months = pd.date_range("2026-09-01", periods=3, freq="MS")
    baseline = pd.DataFrame(
        {
            "site_id": [1, 1, 1, 2, 2, 2],
            "fiscal_month": list(months) * 2,
            "headcount": [100, 99, 98, 50, 49, 48],
            "monthly_attrition": [1.0] * 6,
            "req_fills": [0.5] * 6,
            "labor_usd": [640000.0] * 6,
            "overtime_usd": [32000.0] * 6,
            "other_usd": [0.0] * 6,
            "total_labor_usd": [672000.0] * 6,
        }
    )
    scenario_monthly = {}
    for name in (
        "baseline",
        "attrition_spike",
        "hiring_freeze",
        "overtime_shift",
        "accelerated_hiring",
        "spike_plus_freeze",
    ):
        frame = (
            baseline.groupby("fiscal_month", as_index=False)
            .agg(
                headcount=("headcount", "sum"),
                monthly_attrition=("monthly_attrition", "sum"),
                req_fills=("req_fills", "sum"),
                labor_usd=("labor_usd", "sum"),
                overtime_usd=("overtime_usd", "sum"),
                other_usd=("other_usd", "sum"),
                total_labor_usd=("total_labor_usd", "sum"),
            )
            .round(2)
        )
        scenario_monthly[name] = frame

    budget_variance = pd.DataFrame(
        {
            "site_name": ["Gresham", "Gresham"],
            "cc_code": ["GRS-FAB", "GRS-TNA"],
            "function": ["Fab Operations", "Test & Assembly"],
            "fiscal_month": [months[0], months[0]],
            "planned_total_usd": [1000000.0, 500000.0],
            "actual_total_usd": [1050000.0, 480000.0],
            "total_variance_usd": [50000.0, -20000.0],
            "total_variance_pct": [5.0, -4.0],
            "ytd_variance_usd": [50000.0, -20000.0],
        }
    )
    return {
        "baseline": baseline,
        "blended_salary": {1: 60000.0, 2: 72000.0},
        "sites": pd.DataFrame({"site_id": [1, 2], "site_name": ["Gresham", "Brno"]}),
        "scenario_monthly": scenario_monthly,
        "comparison": pd.DataFrame(),
        "break_even_grid": pd.DataFrame(
            {
                "attrition_multiplier": [1.0, 1.5],
                "backfill_cost_usd": [8.0, 7.0],
                "overtime_policy_cost_usd": [8.2, 7.4],
                "ot_minus_backfill_usd": [0.2, 0.4],
            }
        ),
        "break_even_multiplier": 1.0,
        "break_even_hours": 157.0,
        "budget_variance": budget_variance,
        "cost_centers": pd.DataFrame(
            {
                "cc_code": ["GRS-FAB", "GRS-TNA"],
                "cc_name": ["Gresham Fab Operations", "Gresham Test & Assembly"],
            }
        ),
        "assumptions": {
            "fringe_rate": 0.28,
            "ot_fraction": 0.05,
            "ot_premium": 1.5,
            "attrition_multiplier": 1.75,
            "spike_months": 6,
            "freeze_start_month": 1,
            "max_ot_hours": 60,
            "ot_hours_per_head": 160,
            "hire_cost": 5000,
            "extra_reqs_per_month": 5,
            "sites_count": 2,
            "avg_blended_salary": 64000.0,
        },
        "current_active_total": 150.0,
    }


def test_build_report_creates_valid_workbook(tmp_path):
    path = build_report(_payload(), tmp_path / "model.xlsx")
    assert path.exists()

    wb = load_workbook(path)
    assert wb.sheetnames == ["Assumptions", "Headcount_Forecast", "Scenarios", "Budget_Variance"]

    names = set(wb.defined_names.keys())
    assert {"fringe_rate", "ot_premium", "attrition_multiplier", "hire_cost"} <= names

    hc = wb["Headcount_Forecast"]
    assert str(hc["G2"].value).startswith("=")
    assert hc.freeze_panes == "A2"
    assert hc.auto_filter.ref is not None

    scenarios = wb["Scenarios"]
    assert str(scenarios["F6"].value).startswith("=")
    assert "attrition_multiplier" in str(scenarios["F6"].value)
    assert "fringe_rate" in str(scenarios["J6"].value)
    assert len(scenarios._charts) == 1
    assert len(scenarios.conditional_formatting._cf_rules) >= 2
    assert scenarios.freeze_panes == "A6"

    variance = wb["Budget_Variance"]
    assert str(variance["G2"].value).startswith("=")
    assert len(variance._charts) == 1
    assert variance.freeze_panes == "A2"
    assert variance.auto_filter.ref is not None

    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.endswith((".xml", ".rels")) and name != "[Content_Types].xml":
                ET.fromstring(archive.read(name))


def test_break_even_hours_formula_references_named_ranges(tmp_path):
    path = build_report(_payload(), tmp_path / "model.xlsx")
    wb = load_workbook(path)
    scenarios = wb["Scenarios"]
    formula_cells = [
        cell
        for row in scenarios.iter_rows(min_row=5, max_row=45, min_col=2, max_col=2)
        for cell in row
    ]
    hours_formula = next(
        (
            cell.value
            for cell in formula_cells
            if isinstance(cell.value, str) and "2080" in cell.value
        ),
        None,
    )
    assert hours_formula is not None
    assert "fringe_rate" in hours_formula
    assert "ot_premium" in hours_formula
    assert "hire_cost" in hours_formula
