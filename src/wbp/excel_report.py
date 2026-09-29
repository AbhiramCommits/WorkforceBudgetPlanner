"""Excel deliverable builder (openpyxl): assumptions, named ranges, scenario formulas, charts."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName

MONEY = '"$"#,##0.00'
PCT = "+0.0%;-0.0%;0.0%"
RATE = "0.00"
NUM = "0.0"
INT = "0"

RED_FILL = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")
GREEN_FILL = PatternFill(start_color="FFC6EFCE", end_color="FFC6EFCE", fill_type="solid")
HEADER_FILL = PatternFill(start_color="FFDDEBF7", end_color="FFDDEBF7", fill_type="solid")
TITLE_FONT = Font(bold=True, size=14)
BOLD = Font(bold=True)


def _style_header(ws, row: int, columns: int) -> None:
    for col in range(1, columns + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def _set_widths(ws, widths: dict[str, int]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def _assumptions_sheet(wb: Workbook, payload: dict) -> None:
    ws = wb.active
    ws.title = "Assumptions"
    assumptions = payload["assumptions"]

    ws["A1"] = "Workforce Planning Assumptions"
    ws["A1"].font = TITLE_FONT
    ws.append([])
    ws.append(["Name", "Value", "Unit", "Notes"])
    _style_header(ws, 3, 4)

    rows = [
        ("fringe_rate", assumptions["fringe_rate"], "fraction", "Benefits/load on base pay"),
        (
            "ot_fraction",
            assumptions["ot_fraction"],
            "fraction",
            "Structural overtime share of labor",
        ),
        ("ot_premium", assumptions["ot_premium"], "multiplier", "Overtime pay premium"),
        (
            "attrition_multiplier",
            assumptions["attrition_multiplier"],
            "multiplier",
            "Scenario attrition spike",
        ),
        ("spike_months", assumptions["spike_months"], "months", "Duration of attrition spike"),
        (
            "freeze_start_month",
            assumptions["freeze_start_month"],
            "month",
            "No requisition fills after this month",
        ),
        (
            "max_ot_hours",
            assumptions["max_ot_hours"],
            "hours/head/month",
            "Overtime cap per employee",
        ),
        (
            "ot_hours_per_head",
            assumptions["ot_hours_per_head"],
            "hours/head/month",
            "Monthly capacity of one missing head",
        ),
        ("hire_cost", assumptions["hire_cost"], "USD/hire", "Ramp cost per hire"),
        (
            "extra_reqs_per_month",
            assumptions["extra_reqs_per_month"],
            "reqs/site/month",
            "Accelerated hiring volume",
        ),
        ("sites_count", assumptions["sites_count"], "sites", "Number of sites"),
        (
            "avg_blended_salary",
            assumptions["avg_blended_salary"],
            "USD/year",
            "Headcount-weighted average band salary",
        ),
    ]
    number_formats = {
        "fringe_rate": RATE,
        "ot_fraction": RATE,
        "ot_premium": NUM,
        "attrition_multiplier": RATE,
        "spike_months": INT,
        "freeze_start_month": INT,
        "max_ot_hours": INT,
        "ot_hours_per_head": INT,
        "hire_cost": MONEY,
        "extra_reqs_per_month": INT,
        "sites_count": INT,
        "avg_blended_salary": MONEY,
    }
    for i, (name, value, unit, note) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=name)
        value_cell = ws.cell(row=i, column=2, value=value)
        value_cell.number_format = number_formats[name]
        ws.cell(row=i, column=3, value=unit)
        ws.cell(row=i, column=4, value=note)
        wb.defined_names.add(DefinedName(name, attr_text=f"Assumptions!$B${i}"))

    _set_widths(ws, {"A": 22, "B": 14, "C": 18, "D": 40})
    ws.freeze_panes = "A4"


def _headcount_sheet(wb: Workbook, payload: dict) -> None:
    ws = wb.create_sheet("Headcount_Forecast")
    headers = [
        "Fiscal Month",
        "Site",
        "Headcount",
        "Blended Salary (USD/yr)",
        "Attrition",
        "Req Fills",
        "Labor USD",
        "Overtime USD",
        "Total Labor USD",
    ]
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    baseline = payload["baseline"].sort_values(["site_id", "fiscal_month"])
    site_names = payload["sites"].set_index("site_id")["site_name"].to_dict()
    blended = payload["blended_salary"]
    for _, row in baseline.iterrows():
        r = ws.max_row + 1
        ws.cell(
            row=r, column=1, value=row["fiscal_month"].to_pydatetime()
        ).number_format = "mmm yyyy"
        ws.cell(row=r, column=2, value=site_names[row["site_id"]])
        ws.cell(row=r, column=3, value=row["headcount"]).number_format = NUM
        ws.cell(row=r, column=4, value=blended[row["site_id"]]).number_format = MONEY
        ws.cell(row=r, column=5, value=row["monthly_attrition"]).number_format = RATE
        ws.cell(row=r, column=6, value=row["req_fills"]).number_format = RATE
        ws.cell(
            row=r, column=7, value=f"=ROUND(C{r}*D{r}*(1+fringe_rate)/12,2)"
        ).number_format = MONEY
        ws.cell(row=r, column=8, value=f"=ROUND(G{r}*ot_fraction,2)").number_format = MONEY
        ws.cell(row=r, column=9, value=f"=ROUND(G{r}+H{r},2)").number_format = MONEY

    last = ws.max_row
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{last}"
    _set_widths(
        ws, {"A": 13, "B": 12, "C": 11, "D": 22, "E": 11, "F": 11, "G": 14, "H": 14, "I": 16}
    )


def _scenarios_sheet(wb: Workbook, payload: dict) -> None:
    ws = wb.create_sheet("Scenarios")
    aggregate = payload["scenario_monthly"]

    ws["A1"] = "Scenario Model (live formulas driven by Assumptions named ranges)"
    ws["A1"].font = TITLE_FONT

    hc0_row = 3
    ws.cell(row=hc0_row, column=1, value="Current active headcount (total)")
    ws.cell(row=hc0_row, column=2, value=payload["current_active_total"]).number_format = NUM

    header_row = 5
    headers = [
        "Month",  # A
        "Baseline HC",  # B
        "Baseline Attrition",  # C
        "Baseline Fills",  # D
        "Blended Salary",  # E
        "Spike HC",  # F
        "Freeze HC",  # G
        "OT Shift HC",  # H
        "Accel HC",  # I
        "Baseline Labor",  # J
        "Baseline OT",  # K
        "Spike Labor",  # L
        "Spike OT",  # M
        "OT Shift OT Cost",  # N
        "OT Shift Labor",  # O
        "Accel Labor",  # P
        "Accel OT",  # Q
        "Accel Other",  # R
        "Freeze Labor",  # S
        "Freeze OT",  # T
        "Spike+Freeze HC",  # U
        "Spike+Freeze Labor",  # V
        "Spike+Freeze OT",  # W
    ]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col, value=header)
    _style_header(ws, header_row, len(headers))

    baseline_agg = aggregate["baseline"].set_index("fiscal_month")

    first_row = header_row + 1
    last_row = header_row + len(baseline_agg)
    for offset, month in enumerate(baseline_agg.index):
        r = first_row + offset
        k = offset + 1
        ws.cell(row=r, column=1, value=month.to_pydatetime()).number_format = "mmm yyyy"
        ws.cell(
            row=r, column=2, value=float(baseline_agg.loc[month, "headcount"])
        ).number_format = NUM
        ws.cell(
            row=r, column=3, value=float(baseline_agg.loc[month, "monthly_attrition"])
        ).number_format = RATE
        ws.cell(
            row=r, column=4, value=float(baseline_agg.loc[month, "req_fills"])
        ).number_format = RATE
        blended_t = (
            float(baseline_agg.loc[month, "labor_usd"])
            / max(float(baseline_agg.loc[month, "headcount"]), 1e-9)
            * 12.0
            / (1.0 + payload["assumptions"]["fringe_rate"])
        )
        ws.cell(row=r, column=5, value=blended_t).number_format = MONEY

        spike_cond = f"IF({k}<=spike_months,attrition_multiplier,1)"
        freeze_cond = f"IF({k}>freeze_start_month,0,D{r})"
        spike_prev = "$B$3" if offset == 0 else f"F{r - 1}"
        freeze_prev = "$B$3" if offset == 0 else f"G{r - 1}"
        ot_prev = "$B$3" if offset == 0 else f"H{r - 1}"
        accel_prev = "$B$3" if offset == 0 else f"I{r - 1}"
        composed_prev = "$B$3" if offset == 0 else f"U{r - 1}"

        ws.cell(row=r, column=6, value=f"={spike_prev}-C{r}*{spike_cond}+D{r}").number_format = NUM
        ws.cell(row=r, column=7, value=f"={freeze_prev}-C{r}+{freeze_cond}").number_format = NUM
        ws.cell(row=r, column=8, value=f"={ot_prev}-C{r}").number_format = NUM
        ws.cell(
            row=r, column=9, value=f"={accel_prev}-C{r}+D{r}+extra_reqs_per_month*sites_count"
        ).number_format = NUM
        ws.cell(
            row=r, column=10, value=f"=ROUND(B{r}*E{r}*(1+fringe_rate)/12,2)"
        ).number_format = MONEY
        ws.cell(row=r, column=11, value=f"=ROUND(J{r}*ot_fraction,2)").number_format = MONEY
        ws.cell(
            row=r, column=12, value=f"=ROUND(F{r}*E{r}*(1+fringe_rate)/12,2)"
        ).number_format = MONEY
        ws.cell(row=r, column=13, value=f"=ROUND(L{r}*ot_fraction,2)").number_format = MONEY
        ws.cell(
            row=r,
            column=14,
            value=(
                f"=ROUND(K{r}+MIN(MAX(B{r}-H{r},0)*ot_hours_per_head,H{r}*max_ot_hours)"
                f"*E{r}/2080*ot_premium,2)"
            ),
        ).number_format = MONEY
        ws.cell(
            row=r, column=15, value=f"=ROUND(H{r}*E{r}*(1+fringe_rate)/12,2)"
        ).number_format = MONEY
        ws.cell(
            row=r, column=16, value=f"=ROUND(I{r}*E{r}*(1+fringe_rate)/12,2)"
        ).number_format = MONEY
        ws.cell(row=r, column=17, value=f"=ROUND(P{r}*ot_fraction,2)").number_format = MONEY
        ws.cell(
            row=r, column=18, value="=ROUND(extra_reqs_per_month*sites_count*hire_cost,2)"
        ).number_format = MONEY
        ws.cell(
            row=r, column=19, value=f"=ROUND(G{r}*E{r}*(1+fringe_rate)/12,2)"
        ).number_format = MONEY
        ws.cell(row=r, column=20, value=f"=ROUND(S{r}*ot_fraction,2)").number_format = MONEY
        ws.cell(
            row=r, column=21, value=f"={composed_prev}-C{r}*{spike_cond}+{freeze_cond}"
        ).number_format = NUM
        ws.cell(
            row=r, column=22, value=f"=ROUND(U{r}*E{r}*(1+fringe_rate)/12,2)"
        ).number_format = MONEY
        ws.cell(row=r, column=23, value=f"=ROUND(V{r}*ot_fraction,2)").number_format = MONEY

    comparison_header_row = last_row + 2
    comparison_headers = [
        "Scenario",
        "Ending Headcount",
        "Labor USD",
        "Overtime USD",
        "Other USD",
        "Total USD",
        "Cost per Head",
        "Delta vs Baseline USD",
        "Delta vs Baseline %",
    ]
    for col, header in enumerate(comparison_headers, start=1):
        ws.cell(row=comparison_header_row, column=col, value=header)
    _style_header(ws, comparison_header_row, len(comparison_headers))

    comparison_rows = [
        ("Baseline", "B", "J", "K", None),
        ("Hiring Freeze", "G", "S", "T", None),
        ("Attrition Spike", "F", "L", "M", None),
        ("Overtime Shift", "H", "O", "N", None),
        ("Accelerated Hiring", "I", "P", "Q", "R"),
        ("Spike + Freeze", "U", "V", "W", None),
    ]
    baseline_total_cell = None
    for i, (label, hc_col, labor_col, ot_col, other_col) in enumerate(comparison_rows):
        r = comparison_header_row + 1 + i
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=f"={hc_col}{last_row}").number_format = NUM
        ws.cell(
            row=r, column=3, value=f"=SUM({labor_col}{first_row}:{labor_col}{last_row})"
        ).number_format = MONEY
        ws.cell(
            row=r, column=4, value=f"=SUM({ot_col}{first_row}:{ot_col}{last_row})"
        ).number_format = MONEY
        if other_col:
            ws.cell(
                row=r, column=5, value=f"=SUM({other_col}{first_row}:{other_col}{last_row})"
            ).number_format = MONEY
        else:
            ws.cell(row=r, column=5, value=0).number_format = MONEY
        ws.cell(row=r, column=6, value=f"=C{r}+D{r}+E{r}").number_format = MONEY
        ws.cell(
            row=r, column=7, value=f"=F{r}/AVERAGE({hc_col}{first_row}:{hc_col}{last_row})"
        ).number_format = MONEY
        if i == 0:
            baseline_total_cell = f"$F${r}"
        ws.cell(row=r, column=8, value=f"={baseline_total_cell}-F{r}").number_format = MONEY
        ws.cell(
            row=r, column=9, value=f'=IFERROR(H{r}/{baseline_total_cell},"")'
        ).number_format = PCT

    comp_first = comparison_header_row + 1
    comp_last = comparison_header_row + len(comparison_rows)
    ws.conditional_formatting.add(
        f"H{comp_first}:H{comp_last}",
        FormulaRule(formula=[f"$H{comp_first}<0"], fill=RED_FILL),
    )
    ws.conditional_formatting.add(
        f"H{comp_first}:H{comp_last}",
        FormulaRule(formula=[f"$H{comp_first}>0"], fill=GREEN_FILL),
    )
    ws.conditional_formatting.add(
        f"I{comp_first}:I{comp_last}",
        FormulaRule(formula=[f"$I{comp_first}<-0.05"], fill=RED_FILL),
    )
    ws.conditional_formatting.add(
        f"I{comp_first}:I{comp_last}",
        FormulaRule(formula=[f"$I{comp_first}>0.05"], fill=GREEN_FILL),
    )

    break_even_row = comp_last + 3
    ws.cell(row=break_even_row, column=1, value="Break-even analysis").font = BOLD
    ws.cell(
        row=break_even_row + 1,
        column=1,
        value="Attrition multiplier where overtime policy exceeds backfill hiring:",
    )
    multiplier = payload["break_even_multiplier"]
    ws.cell(
        row=break_even_row + 1,
        column=2,
        value=multiplier if multiplier is not None else "none over tested range",
    ).number_format = RATE
    ws.cell(
        row=break_even_row + 2,
        column=1,
        value="Break-even OT hours per head-month (overtime cost = backfill cost):",
    )
    ws.cell(
        row=break_even_row + 2,
        column=2,
        value="=ROUND(((1+fringe_rate)/12+hire_cost/(12*avg_blended_salary))*2080/ot_premium,0)",
    ).number_format = INT

    grid_header_row = break_even_row + 4
    ws.cell(row=grid_header_row, column=1, value="Attrition Multiplier")
    ws.cell(row=grid_header_row, column=2, value="Backfill Cost USD")
    ws.cell(row=grid_header_row, column=3, value="Overtime Policy Cost USD")
    ws.cell(row=grid_header_row, column=4, value="OT - Backfill USD")
    _style_header(ws, grid_header_row, 4)
    for i, row in payload["break_even_grid"].iterrows():
        r = grid_header_row + 1 + i
        ws.cell(row=r, column=1, value=row["attrition_multiplier"]).number_format = RATE
        ws.cell(row=r, column=2, value=row["backfill_cost_usd"]).number_format = MONEY
        ws.cell(row=r, column=3, value=row["overtime_policy_cost_usd"]).number_format = MONEY
        ws.cell(row=r, column=4, value=row["ot_minus_backfill_usd"]).number_format = MONEY

    ws.cell(
        row=break_even_row + 3,
        column=1,
        value=(
            "Note: OT scenario math is an aggregate approximation; per-site detail "
            "is in data/marts/scenario_monthly.*"
        ),
    )

    chart = LineChart()
    chart.title = "Headcount: Baseline vs Scenarios"
    chart.y_axis.title = "Headcount"
    chart.x_axis.title = "Month"
    chart.height = 9
    chart.width = 22
    chart.add_data(
        Reference(ws, min_col=2, max_col=2, min_row=header_row, max_row=last_row),
        titles_from_data=True,
    )
    for col in (6, 7, 8, 9, 21):
        chart.add_data(
            Reference(ws, min_col=col, max_col=col, min_row=header_row, max_row=last_row),
            titles_from_data=True,
        )
    chart.set_categories(Reference(ws, min_col=1, min_row=first_row, max_row=last_row))
    ws.add_chart(chart, f"A{grid_header_row + len(payload['break_even_grid']) + 3}")

    ws.freeze_panes = "A6"
    _set_widths(
        ws,
        {
            "A": 13,
            "B": 13,
            "C": 13,
            "D": 12,
            "E": 14,
            "F": 11,
            "G": 11,
            "H": 12,
            "I": 11,
            "J": 14,
            "K": 12,
            "L": 12,
            "M": 10,
            "N": 14,
            "O": 12,
            "P": 12,
            "Q": 10,
            "R": 12,
            "S": 12,
            "T": 10,
            "U": 13,
            "V": 14,
            "W": 12,
        },
    )


def _variance_sheet(wb: Workbook, payload: dict) -> None:
    ws = wb.create_sheet("Budget_Variance")
    headers = [
        "Site",
        "Cost Center",
        "Function",
        "Fiscal Month",
        "Planned Total USD",
        "Actual Total USD",
        "Variance USD",
        "Variance %",
        "YTD Variance USD",
    ]
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    for _, row in payload["budget_variance"].iterrows():
        r = ws.max_row + 1
        ws.cell(row=r, column=1, value=row["site_name"])
        ws.cell(row=r, column=2, value=row["cc_code"])
        ws.cell(row=r, column=3, value=row["function"])
        ws.cell(
            row=r, column=4, value=row["fiscal_month"].to_pydatetime()
        ).number_format = "mmm yyyy"
        ws.cell(row=r, column=5, value=float(row["planned_total_usd"])).number_format = MONEY
        ws.cell(row=r, column=6, value=float(row["actual_total_usd"])).number_format = MONEY
        ws.cell(row=r, column=7, value=f"=E{r}-F{r}").number_format = MONEY
        ws.cell(row=r, column=8, value=f'=IFERROR((E{r}-F{r})/E{r},"")').number_format = PCT
        ws.cell(row=r, column=9, value=float(row["ytd_variance_usd"])).number_format = MONEY

    last = ws.max_row
    ws.conditional_formatting.add(f"G2:H{last}", FormulaRule(formula=["$H2<-0.05"], fill=RED_FILL))
    ws.conditional_formatting.add(f"G2:H{last}", FormulaRule(formula=["$H2>0.05"], fill=GREEN_FILL))

    ws.cell(row=1, column=11, value="Cost Center")
    ws.cell(row=1, column=12, value="Variance USD (total)")
    ws.cell(row=1, column=11).font = BOLD
    ws.cell(row=1, column=12).font = BOLD
    ws.cell(row=1, column=11).fill = HEADER_FILL
    ws.cell(row=1, column=12).fill = HEADER_FILL
    for i, row in payload["cost_centers"].iterrows():
        r = 2 + i
        ws.cell(row=r, column=11, value=f'{row["cc_code"]} {row["cc_name"]}')
        ws.cell(
            row=r, column=12, value=f"=SUMIFS($G$2:$G${last},$B$2:$B${last},LEFT(K{r},7))"
        ).number_format = MONEY

    bar = BarChart()
    bar.type = "col"
    bar.grouping = "clustered"
    bar.title = "Variance by Cost Center (USD)"
    bar.height = 10
    bar.width = 22
    bar.add_data(
        Reference(ws, min_col=12, max_col=12, min_row=1, max_row=1 + len(payload["cost_centers"])),
        titles_from_data=True,
    )
    bar.set_categories(
        Reference(ws, min_col=11, min_row=2, max_row=1 + len(payload["cost_centers"]))
    )
    ws.add_chart(bar, "N2")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{last}"
    _set_widths(
        ws,
        {
            "A": 12,
            "B": 11,
            "C": 18,
            "D": 13,
            "E": 17,
            "F": 17,
            "G": 16,
            "H": 12,
            "I": 17,
            "K": 30,
            "L": 22,
        },
    )


def build_report(payload: dict, output_path: str | Path) -> Path:
    wb = Workbook()
    _assumptions_sheet(wb, payload)
    _headcount_sheet(wb, payload)
    _scenarios_sheet(wb, payload)
    _variance_sheet(wb, payload)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    print("build_report is invoked from wbp.scenarios; use `make scenarios`.")
